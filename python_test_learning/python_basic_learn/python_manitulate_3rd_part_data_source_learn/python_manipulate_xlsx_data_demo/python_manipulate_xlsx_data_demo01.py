#!/usr/bin/env python
# -*- encoding: utf-8 -*-
'''
@File    :   python_manipulate_xlsx_data_demo01.py
@Contact :   wangdamar@gmail.com
@License :   GPL v3

@Modify Time      @Author    @Version    @Desciption
------------      -------    --------    -----------
2020/10/11 10:29   banxian      1.0         None
'''

from openpyxl import Workbook
wb = Workbook()

# grab the active worksheet
ws = wb.active

# Data can be assigned directly to cells
ws['A1'] = 42

# Rows can also be appended
ws.append([1, 2, 3])

# Python types will automatically be converted
import datetime
ws['A2'] = datetime.datetime.now()

# Save the file
wb.save("sample.xlsx")